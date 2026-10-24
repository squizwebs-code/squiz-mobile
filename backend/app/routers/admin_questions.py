"""Admin question management, CSV/JSON import (validate -> commit) and export."""
import csv
import hashlib
import io
import json
import re

from fastapi import APIRouter, Depends, File, Form, Query, Request, Response, UploadFile
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import func, or_, select
from sqlalchemy.orm import Session, selectinload

from ..db import get_db
from ..deps import admin_user, audit, err
from ..models import (AnswerOption, Difficulty, Grade, GradeSubject, ImportBatch, PublishStatus, Question, Subject, Unit, User, utcnow)
from ..schemas import OptionOut, Page, QuestionIn, QuestionPatch

router = APIRouter(prefix="/admin/questions", tags=["admin"], dependencies=[Depends(admin_user)])

MAX_IMPORT_BYTES = 5 * 1024 * 1024
MAX_IMPORT_ROWS = 5000
LETTERS = "abcdefgh"


def qkey(text: str) -> str:
    return hashlib.sha256(" ".join(text.split()).casefold().encode()).hexdigest()


def _norm(s: str) -> str:
    return " ".join(str(s).split()).casefold().rstrip(".").strip()


def _nf():
    return err(404, "not_found", "Question not found.")


def admin_out(q: Question, unit: Unit | None = None, subject: Subject | None = None, grade: Grade | None = None) -> dict:
    return {"id": q.id, "unit_id": q.unit_id, "unit_title": unit.title if unit else None, "subject_id": unit.subject_id if unit else None,
            "subject": subject.name if subject else None, "grade_id": unit.grade_id if unit else None, "grade_number": grade.number if grade else None,
            "text": q.text, "explanation": q.explanation, "difficulty": q.difficulty.value if q.difficulty else None, "image_url": q.image_url,
            "status": q.status.value, "version": q.version, "created_at": q.created_at, "updated_at": q.updated_at, "deleted": q.deleted_at is not None,
            "options": [{"id": o.id, "text": o.text, "is_correct": o.is_correct, "position": o.position} for o in q.options]}


def _full_out(db: Session, q: Question) -> dict:
    u = db.get(Unit, q.unit_id)
    return admin_out(q, u, db.get(Subject, u.subject_id), db.get(Grade, u.grade_id))


# ------------------------------------------------------------------ list / CRUD
@router.get("")
def list_questions(q: str | None = None, status: str | None = None, grade_id: int | None = None, subject_id: int | None = None, unit_id: int | None = None,
                   difficulty: str | None = None, include_deleted: bool = False, sort: str = "updated", order: str = "desc",
                   page: int = Query(1, ge=1), size: int = Query(20, ge=1, le=100), db: Session = Depends(get_db)):
    st = select(Question).join(Unit, Unit.id == Question.unit_id)
    if not include_deleted:
        st = st.where(Question.deleted_at.is_(None))
    if q:
        st = st.where(Question.text.ilike(f"%{q.strip().replace('%', '').replace('_', '')}%"))
    if status:
        st = st.where(Question.status == PublishStatus(status))
    if grade_id:
        st = st.where(Unit.grade_id == grade_id)
    if subject_id:
        st = st.where(Unit.subject_id == subject_id)
    if unit_id:
        st = st.where(Question.unit_id == unit_id)
    if difficulty:
        st = st.where(Question.difficulty == Difficulty(difficulty))
    col = {"updated": Question.updated_at, "created": Question.created_at, "text": Question.text, "id": Question.id}.get(sort, Question.updated_at)
    st = st.order_by(col.asc() if order == "asc" else col.desc(), Question.id.desc())
    total = db.scalar(select(func.count()).select_from(st.subquery()))
    rows = db.scalars(st.options(selectinload(Question.options)).offset((page - 1) * size).limit(size)).all()
    units = {u.id: u for u in db.scalars(select(Unit).where(Unit.id.in_({r.unit_id for r in rows} or {0})))}
    subs = {s.id: s for s in db.scalars(select(Subject))}
    grades = {g.id: g for g in db.scalars(select(Grade))}
    items = [admin_out(r, units[r.unit_id], subs[units[r.unit_id].subject_id], grades[units[r.unit_id].grade_id]) for r in rows]
    return Page(items=items, total=total, page=page, size=size)


def _check_unit(db: Session, unit_id: int) -> Unit:
    u = db.get(Unit, unit_id)
    if not u:
        raise err(422, "unit_not_found", "The selected unit does not exist.")
    return u


@router.post("", status_code=201)
def create_question(body: QuestionIn, request: Request, db: Session = Depends(get_db), admin: User = Depends(admin_user)):
    _check_unit(db, body.unit_id)
    key = qkey(body.text)
    if db.scalar(select(Question.id).where(Question.unit_id == body.unit_id, Question.text_key == key, Question.deleted_at.is_(None))):
        raise err(409, "duplicate_question", "This unit already contains a question with the same text.")
    q = Question(unit_id=body.unit_id, text=body.text, text_key=key, explanation=(body.explanation or "").strip() or None, difficulty=body.difficulty,
                 image_url=body.image_url, status=body.status)
    q.options = [AnswerOption(text=o.text, is_correct=o.is_correct, position=i) for i, o in enumerate(body.options)]
    db.add(q)
    db.flush()
    audit(db, admin, "question.create", "question", q.id, {"unit_id": q.unit_id, "status": q.status.value}, request)
    db.commit()
    return _full_out(db, q)


def _get(db: Session, qid: int) -> Question:
    q = db.scalar(select(Question).where(Question.id == qid).options(selectinload(Question.options)))
    if not q:
        raise _nf()
    return q


@router.get("/export")
def export_questions(format: str = Query("json", pattern="^(json|csv|xlsx)$"), status: str | None = None, grade_id: int | None = None,
                     subject_id: int | None = None, unit_id: int | None = None, request: Request = None, db: Session = Depends(get_db),
                     admin: User = Depends(admin_user)):
    st = select(Question).join(Unit, Unit.id == Question.unit_id).where(Question.deleted_at.is_(None))
    if status:
        st = st.where(Question.status == PublishStatus(status))
    if grade_id:
        st = st.where(Unit.grade_id == grade_id)
    if subject_id:
        st = st.where(Unit.subject_id == subject_id)
    if unit_id:
        st = st.where(Question.unit_id == unit_id)
    rows = db.scalars(st.options(selectinload(Question.options)).order_by(Unit.grade_id, Unit.subject_id, Unit.position, Question.id)).all()
    units = {u.id: u for u in db.scalars(select(Unit))}
    subs = {s.id: s for s in db.scalars(select(Subject))}
    grades = {g.id: g for g in db.scalars(select(Grade))}
    recs = []
    for q in rows:
        u = units[q.unit_id]
        recs.append({"grade": grades[u.grade_id].number, "subject": subs[u.subject_id].key, "unit": u.title, "question": q.text,
                     "options": [{"text": o.text, "is_correct": o.is_correct} for o in q.options], "explanation": q.explanation or "",
                     "image": q.image_url or "", "difficulty": q.difficulty.value if q.difficulty else "", "status": q.status.value})
    audit(db, admin, "question.export", "question", None, {"count": len(recs), "format": format}, request)
    db.commit()
    if format == "xlsx":
        wb = Workbook()
        ws = wb.active
        ws.title = "Questions"
        ws.append(TEMPLATE_COLUMNS + ["status"])
        for r in recs:
            opts = [o["text"] for o in r["options"]][:4]
            letter = next((LETTERS[i].upper() for i, o in enumerate(r["options"][:4]) if o["is_correct"]), "")
            ws.append([_csv_safe(x) for x in [r["grade"], r["subject"], r["unit"], r["question"]] + opts + [""] * (4 - len(opts)) +
                       [letter, r["explanation"], r.get("image") or "", r["difficulty"], r["status"]]])
        _style_header(ws)
        out = io.BytesIO()
        wb.save(out)
        return Response(out.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        headers={"Content-Disposition": 'attachment; filename="squiz-questions.xlsx"'})
    if format == "json":
        return Response(json.dumps({"questions": recs}, ensure_ascii=False, indent=1), media_type="application/json",
                        headers={"Content-Disposition": 'attachment; filename="squiz-questions.json"'})
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["grade", "subject", "unit", "question"] + [f"option_{c}" for c in LETTERS] + ["answer", "explanation", "difficulty", "status"])
    for r in recs:
        opts = [o["text"] for o in r["options"]]
        ans = next((o["text"] for o in r["options"] if o["is_correct"]), "")
        w.writerow([_csv_safe(x) for x in [r["grade"], r["subject"], r["unit"], r["question"]] + opts + [""] * (8 - len(opts)) + [ans, r["explanation"], r["difficulty"], r["status"]]])
    return Response("﻿" + buf.getvalue(), media_type="text/csv; charset=utf-8", headers={"Content-Disposition": 'attachment; filename="squiz-questions.csv"'})


def _csv_safe(v):
    s = str(v)
    return "'" + s if s[:1] in ("=", "+", "-", "@") and not re.fullmatch(r"-?\d+(\.\d+)?", s) else s     # spreadsheet formula-injection guard


def _csv_unsafe(s: str) -> str:
    return s[1:] if len(s) > 1 and s[0] == "'" and s[1] in ("=", "+", "-", "@") else s


# ------------------------------------------------------------------ import
TEMPLATE_COLUMNS = ["grade", "subject", "unit", "question", "option_a", "option_b", "option_c", "option_d", "correct_answer", "explanation", "image", "difficulty"]
_ALIASES = {"correct": "answer", "correct_answer": "answer", "correct_option": "answer", "answer_key": "answer", "image_url": "image", "text": "question",
            **{f"option {c}": f"option_{c}" for c in LETTERS}, **{c: f"option_{c}" for c in "abcd"}}


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return _csv_unsafe(" ".join(str(v).split())) if not isinstance(v, str) else _csv_unsafe(v.strip())


def _parse_xlsx(data: bytes) -> list[dict]:
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        raise err(422, "bad_xlsx", "This is not a valid Excel (.xlsx) file.")
    ws = wb["Questions"] if "Questions" in wb.sheetnames else wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    head = next(it, None)
    if not head:
        raise err(422, "no_rows", "The workbook is empty.")
    cols = [_ALIASES.get(str(h or "").strip().lower().replace(" ", "_"), _ALIASES.get(str(h or "").strip().lower(), str(h or "").strip().lower().replace(" ", "_"))) for h in head]
    missing = [c for c in ("grade", "subject", "unit", "question", "answer") if c not in cols]
    if missing or not any(c.startswith("option_") for c in cols):
        raise err(422, "bad_columns", "Missing columns: " + ", ".join(missing or ["option_a, option_b"]) + ". Download the template to see the expected columns.")
    rows = []
    for raw in it:
        r = {c: _cell(v) for c, v in zip(cols, raw)}
        if not any(r.values()):
            continue                                           # blank line
        if len(rows) > MAX_IMPORT_ROWS:
            break
        opts = [o for o in (r.get(f"option_{c}", "") for c in LETTERS) if o]
        rows.append({"grade": r.get("grade"), "subject": r.get("subject"), "unit": r.get("unit"), "question": r.get("question"), "options": opts,
                     "answer": r.get("answer"), "explanation": r.get("explanation"), "difficulty": r.get("difficulty"), "status": r.get("status"), "image": r.get("image")})
    return rows


def _parse(filename: str, data: bytes) -> list[dict]:
    if len(data) > MAX_IMPORT_BYTES:
        raise err(413, "file_too_large", "Import files can be at most 5 MB.")
    if filename.lower().endswith((".xlsx", ".xlsm")) or data[:2] == b"PK":
        rows = _parse_xlsx(data)
        if len(rows) > MAX_IMPORT_ROWS:
            raise err(413, "too_many_rows", f"Import at most {MAX_IMPORT_ROWS} questions at a time.")
        if not rows:
            raise err(422, "no_rows", "The file contains no questions.")
        return rows
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        raise err(422, "bad_encoding", "The file must be UTF-8 encoded.")
    rows: list[dict] = []
    if filename.lower().endswith(".json") or text.lstrip().startswith(("[", "{")):
        try:
            obj = json.loads(text)
        except json.JSONDecodeError as e:
            raise err(422, "bad_json", f"Invalid JSON: {e.msg} (line {e.lineno}).")
        items = obj.get("questions") if isinstance(obj, dict) else obj
        if not isinstance(items, list):
            raise err(422, "bad_json", "Expected a list of questions, or an object with a 'questions' list.")
        for it in items:
            if not isinstance(it, dict):
                rows.append({"_bad": True})
                continue
            opts, flagged = [], None
            for i, o in enumerate(it.get("options") or []):
                if isinstance(o, dict):
                    opts.append(str(o.get("text", "")))
                    if o.get("is_correct") or o.get("correct"):
                        flagged = i if flagged is None else -1
                else:
                    opts.append(str(o))
            ans = it.get("answer")
            if ans is None and flagged is not None and flagged >= 0:
                ans = opts[flagged]
            rows.append({"grade": it.get("grade"), "subject": it.get("subject"), "unit": it.get("unit"), "question": it.get("question") or it.get("text"),
                         "options": opts, "answer": ans, "explanation": it.get("explanation"), "difficulty": it.get("difficulty"), "status": it.get("status"),
                         "_multi_correct": flagged == -1})
    else:
        rd = csv.DictReader(io.StringIO(text))
        if not rd.fieldnames:
            raise err(422, "bad_csv", "The CSV file is empty.")
        for r in rd:
            r = {(k or "").strip().lower(): (v or "").strip() for k, v in r.items()}
            opts = [_csv_unsafe(r.get(f"option_{c}", "")) for c in LETTERS]
            opts = [o for o in opts if o]
            rows.append({"grade": r.get("grade"), "subject": _csv_unsafe(r.get("subject", "")), "unit": _csv_unsafe(r.get("unit", "")),
                         "question": _csv_unsafe(r.get("question", "")), "options": opts, "answer": _csv_unsafe(r.get("answer", "")),
                         "explanation": _csv_unsafe(r.get("explanation", "")), "difficulty": r.get("difficulty"), "status": r.get("status")})
    if len(rows) > MAX_IMPORT_ROWS:
        raise err(413, "too_many_rows", f"Import at most {MAX_IMPORT_ROWS} questions at a time.")
    if not rows:
        raise err(422, "no_rows", "The file contains no questions.")
    return rows


def _slug(name: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", name.lower()).strip("-")[:40] or "subject"


def validate_rows(db: Session, rows: list[dict], create_missing: bool) -> list[dict]:
    grades = {g.number: g for g in db.scalars(select(Grade))}
    subs = list(db.scalars(select(Subject)))
    units = list(db.scalars(select(Unit)))
    existing = {(k, u): qid for k, u, qid in db.execute(select(Question.text_key, Question.unit_id, Question.id).where(Question.deleted_at.is_(None)).order_by(Question.id.desc()))}
    new_subjects: dict[str, str] = {}            # slug -> name (planned)
    new_units: dict[tuple, str] = {}              # (grade_no, subject_slug, unit_norm) -> title
    seen: set[tuple] = set()
    results = []
    for i, r in enumerate(rows, start=1):
        errors: list[str] = []
        res = {"row": i, "status": "ok", "errors": errors, "question": (str(r.get("question") or ""))[:90]}
        if r.get("_bad"):
            results.append({**res, "status": "error", "errors": ["Row is not an object"]})
            continue
        # grade
        try:
            gno = int(str(r.get("grade")).strip().lower().replace("grade", "").strip())
        except (TypeError, ValueError):
            gno = None
            errors.append("Grade must be a number such as 9, 10, 11 or 12")
        grade = grades.get(gno) if gno is not None else None
        if gno is not None and grade is None:
            errors.append(f"Grade {gno} does not exist")
        # subject
        sname = str(r.get("subject") or "").strip()
        subject = next((s for s in subs if _norm(s.key) == _norm(sname) or _norm(s.name) == _norm(sname)), None) if sname else None
        sslug = _slug(sname)
        if not sname:
            errors.append("Subject is required")
        elif subject is None:
            if create_missing:
                new_subjects.setdefault(sslug, sname)
            else:
                errors.append(f'Unknown subject "{sname}" (enable "create missing subjects and units" to add it)')
        # unit
        uname = str(r.get("unit") or "").strip()
        if not uname:
            errors.append("Unit is required")
        unit = None
        if uname and grade and subject:
            cand = [u for u in units if u.grade_id == grade.id and u.subject_id == subject.id]
            m = re.fullmatch(r"(?:unit[\s-]*)?(\d+)", uname, re.I)
            wanted = {_norm(uname)} | ({_norm(f"Unit {m.group(1)}"), _norm(f"Unit-{m.group(1)}")} if m else set())
            unit = next((u for u in cand if _norm(u.title) in wanted), None)
        if uname and unit is None and not errors:
            if create_missing:
                title = f"Unit {m.group(1)}" if (m := re.fullmatch(r"(?:unit[\s-]*)?(\d+)", uname, re.I)) else uname
                new_units.setdefault((gno, sslug if subject is None else subject.key, _norm(title)), title)
            else:
                errors.append(f'Unknown unit "{uname}" for this grade and subject')
        # question text and options
        text = " ".join(str(r.get("question") or "").split())
        if len(text) < 3:
            errors.append("Question text is required")
        opts = [" ".join(str(o).split()) for o in (r.get("options") or []) if str(o).strip()]
        if not 2 <= len(opts) <= 8:
            errors.append("Provide between 2 and 8 answer options")
        if len({_norm(o) for o in opts}) != len(opts):
            errors.append("Answer options must be different from each other")
        correct = -1
        if r.get("_multi_correct"):
            errors.append("More than one option is marked correct")
        elif opts:
            ans = str(r.get("answer") or "").strip()
            if not ans:
                errors.append("Correct answer is required")
            else:
                hit = [k for k, o in enumerate(opts) if _norm(o) == _norm(ans)]
                if not hit and len(ans) == 1 and ans.lower() in LETTERS[:len(opts)]:
                    hit = [LETTERS.index(ans.lower())]
                if len(hit) == 1:
                    correct = hit[0]
                else:
                    errors.append("Correct answer does not match any option" if not hit else "Correct answer matches more than one option")
        image = str(r.get("image") or "").strip()
        if image and not image.startswith("/media/"):
            errors.append("Image must be a file uploaded in the admin panel (it starts with /media/)")
        diff = str(r.get("difficulty") or "").strip().lower()
        if diff and diff not in ("easy", "medium", "hard"):
            errors.append("Difficulty must be easy, medium or hard")
        status = str(r.get("status") or "draft").strip().lower() or "draft"
        if status not in ("draft", "published"):
            errors.append("Status must be draft or published")
        if not errors:
            unit_key = unit.id if unit else (gno, sslug if subject is None else subject.key, _norm(f"Unit {m.group(1)}" if (m := re.fullmatch(r"(?:unit[\s-]*)?(\d+)", uname, re.I)) else uname))
            k = (qkey(text), unit_key)
            plan = {"grade": gno, "subject_id": subject.id if subject else None, "subject_name": sname, "subject_slug": sslug,
                    "unit_id": unit.id if unit else None, "unit_title": None if unit else (f"Unit {m.group(1)}" if (m := re.fullmatch(r"(?:unit[\s-]*)?(\d+)", uname, re.I)) else uname),
                    "text": text, "options": opts, "correct": correct, "explanation": " ".join(str(r.get("explanation") or "").split()) or None,
                    "difficulty": diff or None, "status": status, "image": image or None}
            if k in seen:
                res["status"] = "duplicate"
                errors.append("Duplicate of an earlier row in this file")
            elif unit and (k[0], unit.id) in existing:
                res["status"] = "duplicate"
                errors.append("Same question text already exists in this unit")
                res["_plan"] = {**plan, "existing_id": existing[(k[0], unit.id)]}
                seen.add(k)
            else:
                seen.add(k)
                res["_plan"] = plan
        else:
            res["status"] = "error"
        results.append(res)
    return results


def _summary(results: list[dict]) -> dict:
    return {"total": len(results), "ok": sum(r["status"] == "ok" for r in results), "errors": sum(r["status"] == "error" for r in results),
            "duplicates": sum(r["status"] == "duplicate" for r in results)}


def _public(results: list[dict]) -> list[dict]:
    return [{k: v for k, v in r.items() if k != "_plan"} for r in results]


def _style_header(ws) -> None:
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="428BCA")
    ws.freeze_panes = "A2"
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = min(50, max(12, max(len(str(c.value or "")) for c in col[:50]) + 2))


@router.get("/import/template")
def import_template(db: Session = Depends(get_db)):
    """Fixed-column Excel template with an instructions sheet and a reference sheet of the grades, subjects and units that exist now."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Questions"
    ws.append(TEMPLATE_COLUMNS)
    ws.append([12, "physics", "Unit 1", "What is the SI unit of force?", "Joule", "Newton", "Watt", "Pascal", "B", "Force is measured in newtons (N).", "", "easy"])
    ws.append([12, "history", "Unit 2", "Example question text goes here", "Option A", "Option B", "Option C", "Option D", "Option C", "", "", ""])
    _style_header(ws)
    ins = wb.create_sheet("Instructions")
    for line in ["How to fill this sheet", "",
                 "1. One question per row. Do not rename or reorder the column headings in the first row of the 'Questions' sheet.",
                 "2. grade: 9, 10, 11 or 12.  subject: the subject key or name (see the 'Reference' sheet).  unit: 'Unit 3' or just 3, or the unit title.",
                 "3. option_a to option_d are the answers (at least two). You may add option_e to option_h.",
                 "4. correct_answer: the letter (A, B, C, D) or the full text of the right option.",
                 "5. explanation, image and difficulty (easy, medium, hard) are optional. An image must first be uploaded in the admin panel; paste its /media/... path.",
                 "6. Upload the file in Import. You will see a preview with row errors and duplicates before anything is saved.",
                 "7. Imported questions are saved as DRAFTS. Review and publish them when you are ready.",
                 "Delete the two example rows before importing real data."]:
        ins.append([line])
    ins.column_dimensions["A"].width = 140
    ins["A1"].font = Font(bold=True, size=14)
    ref = wb.create_sheet("Reference")
    ref.append(["grade", "subject_key", "subject_name", "track", "unit"])
    grades = {g.id: g.number for g in db.scalars(select(Grade))}
    subs = {x.id: x for x in db.scalars(select(Subject).where(Subject.archived_at.is_(None)))}
    for u in db.scalars(select(Unit).where(Unit.archived_at.is_(None)).order_by(Unit.grade_id, Unit.subject_id, Unit.position)):
        sb = subs.get(u.subject_id)
        if sb:
            ref.append([grades[u.grade_id], sb.key, sb.name, sb.track, u.title])
    _style_header(ref)
    out = io.BytesIO()
    wb.save(out)
    return Response(out.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": 'attachment; filename="squiz-question-template.xlsx"'})


@router.post("/import/validate")
async def import_validate(file: UploadFile = File(...), create_missing: bool = Form(False), db: Session = Depends(get_db)):
    rows = _parse(file.filename or "upload", await file.read(MAX_IMPORT_BYTES + 1))
    results = validate_rows(db, rows, create_missing)
    return {"summary": _summary(results), "rows": _public(results)[:500]}


def _snapshot(q: Question) -> dict:
    return {"text": q.text, "explanation": q.explanation, "difficulty": q.difficulty.value if q.difficulty else None, "image_url": q.image_url,
            "options": [{"text": o.text, "is_correct": o.is_correct} for o in q.options]}


def _apply_options(q: Question, opts: list[tuple[str, bool]]) -> None:
    old = list(q.options)
    for i, (text, ok) in enumerate(opts):
        if i < len(old):
            old[i].text, old[i].is_correct, old[i].position = text, ok, i
        else:
            q.options.append(AnswerOption(text=text, is_correct=ok, position=i))
    for extra in old[len(opts):]:
        q.options.remove(extra)


@router.post("/import/commit")
async def import_commit(request: Request, file: UploadFile = File(...), create_missing: bool = Form(False),
                        policy: str = Form("skip", pattern="^(skip|update|add)$"), publish: bool = Form(False),
                        db: Session = Depends(get_db), admin: User = Depends(admin_user)):
    """Re-validates server-side (never trusts an earlier preview) and imports in one transaction.

    policy for questions that already exist in the unit: skip (default), update them, or add a second copy.
    New questions are DRAFTS unless publish=true, which honours each row's status column. Everything is recorded
    as a batch so it can be undone."""
    rows = _parse(file.filename or "upload", await file.read(MAX_IMPORT_BYTES + 1))
    results = validate_rows(db, rows, create_missing)
    counts = {"subjects": 0, "units": 0, "created": 0, "updated": 0, "skipped_duplicates": 0}
    items: list[dict] = []
    subj_cache: dict[str, Subject] = {_norm(s.key): s for s in db.scalars(select(Subject))}
    grades = {g.number: g for g in db.scalars(select(Grade))}
    unit_cache: dict[tuple, Unit] = {}
    try:
        for r in results:
            p = r.get("_plan")
            if not p or r["status"] not in ("ok", "duplicate"):
                continue
            if r["status"] == "duplicate":
                if not p.get("existing_id") or policy == "skip":
                    counts["skipped_duplicates"] += 1
                    continue
                if policy == "update":
                    q = _get(db, p["existing_id"])
                    items.append({"op": "update", "qid": q.id, "before": _snapshot(q)})
                    _apply_options(q, [(t, k == p["correct"]) for k, t in enumerate(p["options"])])
                    q.explanation = p["explanation"] or q.explanation
                    q.difficulty = Difficulty(p["difficulty"]) if p["difficulty"] else q.difficulty
                    q.image_url = p["image"] or q.image_url
                    q.version += 1
                    counts["updated"] += 1
                    continue
            grade = grades[p["grade"]]
            subject = db.get(Subject, p["subject_id"]) if p["subject_id"] else subj_cache.get(_norm(p["subject_slug"]))
            if subject is None:
                subject = Subject(key=p["subject_slug"], name=p["subject_name"], published=True)
                db.add(subject)
                db.flush()
                subj_cache[_norm(subject.key)] = subject
                counts["subjects"] += 1
                items.append({"op": "subject", "id": subject.id})
            if db.get(GradeSubject, (grade.id, subject.id)) is None:
                db.add(GradeSubject(grade_id=grade.id, subject_id=subject.id, published=True))
                db.flush()
            if p["unit_id"]:
                unit_id = p["unit_id"]
            else:
                key = (grade.id, subject.id, _norm(p["unit_title"]))
                unit = unit_cache.get(key) or db.scalar(select(Unit).where(Unit.grade_id == grade.id, Unit.subject_id == subject.id, func.lower(Unit.title) == p["unit_title"].lower()))
                if unit is None:
                    pos = (db.scalar(select(func.max(Unit.position)).where(Unit.grade_id == grade.id, Unit.subject_id == subject.id)) or 0) + 1
                    unit = Unit(grade_id=grade.id, subject_id=subject.id, title=p["unit_title"], position=pos, published=True)
                    db.add(unit)
                    db.flush()
                    counts["units"] += 1
                    items.append({"op": "unit", "id": unit.id})
                unit_cache[key] = unit
                unit_id = unit.id
            q = Question(unit_id=unit_id, text=p["text"], text_key=qkey(p["text"]), explanation=p["explanation"], image_url=p["image"],
                         difficulty=Difficulty(p["difficulty"]) if p["difficulty"] else None,
                         status=PublishStatus(p["status"]) if publish else PublishStatus.draft, source_ref=f"import:{file.filename}#row{r['row']}"[:160])
            q.options = [AnswerOption(text=t, is_correct=(k == p["correct"]), position=k) for k, t in enumerate(p["options"])]
            db.add(q)
            db.flush()
            items.append({"op": "create", "qid": q.id})
            counts["created"] += 1
        summary = _summary(results)
        batch = ImportBatch(admin_id=admin.id, admin_name=admin.full_name, filename=(file.filename or "upload")[:200], policy=policy,
                            counts={**counts, **summary}, items=items)
        db.add(batch)
        db.flush()
        audit(db, admin, "question.import", "import_batch", batch.id, {"file": file.filename, "policy": policy, **counts, **summary}, request)
        db.commit()
    except Exception:
        db.rollback()
        raise
    return {"batch_id": batch.id, "summary": summary, "created": counts, "rows": _public(results)[:500]}


@router.get("/import/history")
def import_history(page: int = Query(1, ge=1), size: int = Query(20, ge=1, le=100), db: Session = Depends(get_db)):
    st = select(ImportBatch)
    total = db.scalar(select(func.count()).select_from(st.subquery()))
    rows = db.scalars(st.order_by(ImportBatch.id.desc()).offset((page - 1) * size).limit(size)).all()
    return Page(items=[{"id": b.id, "filename": b.filename, "admin": b.admin_name, "policy": b.policy, "status": b.status, "counts": b.counts,
                        "created_at": b.created_at, "undone_at": b.undone_at} for b in rows], total=total, page=page, size=size)


@router.post("/import/{batch_id:int}/undo")
def import_undo(batch_id: int, request: Request, db: Session = Depends(get_db), admin: User = Depends(admin_user)):
    """Reverses a whole batch: new questions are removed (students' past results keep their own snapshot), updated ones are restored,
    and subjects/units the batch created are archived when nothing else uses them."""
    b = db.get(ImportBatch, batch_id)
    if not b:
        raise err(404, "not_found", "Import not found.")
    if b.status == "undone":
        raise err(409, "already_undone", "This import was already undone.")
    removed = restored = 0
    for it in reversed(b.items):
        if it["op"] == "create":
            q = db.get(Question, it["qid"])
            if q and q.deleted_at is None:
                q.deleted_at, q.status = utcnow(), PublishStatus.archived
                removed += 1
        elif it["op"] == "update":
            q = _get(db, it["qid"])
            bf = it["before"]
            q.text, q.text_key, q.explanation, q.image_url = bf["text"], qkey(bf["text"]), bf["explanation"], bf["image_url"]
            q.difficulty = Difficulty(bf["difficulty"]) if bf["difficulty"] else None
            _apply_options(q, [(o["text"], o["is_correct"]) for o in bf["options"]])
            q.version += 1
            restored += 1
        elif it["op"] == "unit":
            db.flush()
            u = db.get(Unit, it["id"])
            if u and not db.scalar(select(func.count()).select_from(Question).where(Question.unit_id == u.id, Question.deleted_at.is_(None))):
                u.archived_at = utcnow()
        elif it["op"] == "subject":
            db.flush()
            sb = db.get(Subject, it["id"])
            if sb and not db.scalar(select(func.count()).select_from(Unit).where(Unit.subject_id == sb.id, Unit.archived_at.is_(None))):
                sb.archived_at = utcnow()
    b.status, b.undone_at = "undone", utcnow()
    audit(db, admin, "question.import_undo", "import_batch", b.id, {"removed": removed, "restored": restored}, request)
    db.commit()
    return {"removed": removed, "restored": restored}


# ------------------------------------------------------------------ item routes (after the fixed paths above)
@router.get("/{question_id:int}")
def get_question(question_id: int, db: Session = Depends(get_db)):
    return _full_out(db, _get(db, question_id))


@router.get("/{question_id:int}/preview")
def preview(question_id: int, db: Session = Depends(get_db)):
    """Shows the question exactly as a student would see it (no correctness info in `student_view`)."""
    q = _get(db, question_id)
    correct = next((o.id for o in q.options if o.is_correct), None)
    return {"student_view": {"id": q.id, "text": q.text, "image_url": q.image_url, "difficulty": q.difficulty.value if q.difficulty else None,
                             "options": [OptionOut(id=o.id, text=o.text).model_dump() for o in q.options]},
            "correct_option_id": correct, "explanation": q.explanation, "status": q.status.value}


@router.patch("/{question_id:int}")
def update_question(question_id: int, body: QuestionPatch, request: Request, db: Session = Depends(get_db), admin: User = Depends(admin_user)):
    q = _get(db, question_id)
    changed = []
    if body.unit_id is not None and body.unit_id != q.unit_id:
        _check_unit(db, body.unit_id)
        q.unit_id = body.unit_id
        changed.append("unit")
    if body.text is not None and " ".join(body.text.split()) != q.text:
        q.text, q.text_key = " ".join(body.text.split()), qkey(body.text)
        changed.append("text")
    if (body.text is not None or body.unit_id is not None) and db.scalar(select(Question.id).where(Question.unit_id == q.unit_id, Question.text_key == q.text_key, Question.deleted_at.is_(None), Question.id != q.id)):
        db.rollback()
        raise err(409, "duplicate_question", "This unit already contains a question with the same text.")
    if body.explanation is not None:
        q.explanation = body.explanation.strip() or None
        changed.append("explanation")
    if body.clear_explanation:
        q.explanation = None
    if body.difficulty is not None:
        q.difficulty = body.difficulty
    if body.clear_difficulty:
        q.difficulty = None
    if body.image_url is not None:
        q.image_url = body.image_url
    if body.clear_image:
        q.image_url = None
    if body.options is not None:
        old = list(q.options)
        for i, o in enumerate(body.options):
            if i < len(old):
                old[i].text, old[i].is_correct, old[i].position = o.text, o.is_correct, i      # keep ids so caches stay valid
            else:
                q.options.append(AnswerOption(text=o.text, is_correct=o.is_correct, position=i))
        for extra in old[len(body.options):]:
            q.options.remove(extra)
        changed.append("options")
    if {"text", "options", "unit"} & set(changed):
        q.version += 1          # past quiz results keep their own snapshot, so history is unaffected
    audit(db, admin, "question.update", "question", q.id, {"changed": changed, "version": q.version}, request)
    db.commit()
    return _full_out(db, _get(db, q.id))


def _set_status(db: Session, request: Request, admin: User, qid: int, status: PublishStatus, action: str):
    q = _get(db, qid)
    if q.deleted_at is not None:
        raise err(409, "deleted", "This question was deleted.")
    q.status = status
    audit(db, admin, action, "question", q.id, None, request)
    db.commit()
    return _full_out(db, q)


@router.post("/{question_id:int}/publish")
def publish(question_id: int, request: Request, db: Session = Depends(get_db), admin: User = Depends(admin_user)):
    return _set_status(db, request, admin, question_id, PublishStatus.published, "question.publish")


@router.post("/{question_id:int}/unpublish")
def unpublish(question_id: int, request: Request, db: Session = Depends(get_db), admin: User = Depends(admin_user)):
    return _set_status(db, request, admin, question_id, PublishStatus.draft, "question.unpublish")


@router.post("/{question_id:int}/archive")
def archive(question_id: int, request: Request, db: Session = Depends(get_db), admin: User = Depends(admin_user)):
    return _set_status(db, request, admin, question_id, PublishStatus.archived, "question.archive")


@router.delete("/{question_id:int}", status_code=204)
def delete_question(question_id: int, request: Request, confirm: bool = Query(False), db: Session = Depends(get_db), admin: User = Depends(admin_user)):
    """Soft delete: hidden everywhere, but students' past results keep their own snapshot of the question."""
    if not confirm:
        raise err(400, "confirmation_required", "Add confirm=true to delete this question.")
    q = _get(db, question_id)
    q.deleted_at = func.now()
    q.status = PublishStatus.archived
    audit(db, admin, "question.delete", "question", q.id, {"text": q.text[:80]}, request)
    db.commit()
    return Response(status_code=204)
