"""Admin features: user management with the promotion chain, suspend, tracks, archiving, Excel import/undo, statistics, announcements, reports."""
import io
import time
from datetime import timedelta

import pytest
from openpyxl import Workbook, load_workbook

from app.models import User, utcnow
from app.routers.admin_questions import qkey
from tests.conftest import ADMIN_PW, STUDENT_PW, _h, make_content, register


def login(client, ident, pw=STUDENT_PW):
    r = client.post("/auth/login", json={"identifier": ident, "password": pw})
    assert r.status_code == 200, r.text
    return _h(r.json()["access_token"]), r.json()


def mk(client, name, phone, username):
    t = register(client, name=name, phone=phone, username=username)
    return t["user"]["id"], _h(t["access_token"])


def promote(client, actor_h, uid):
    r = client.post(f"/admin/users/{uid}/role", json={"role": "admin"}, headers=actor_h)
    assert r.status_code == 200, r.text
    return r.json()


# ------------------------------------------------------------------ users list
def test_admin_sees_every_account_with_details(client, admin_h, student):
    r = client.get("/admin/users", headers=admin_h)
    assert r.status_code == 200
    body = r.json()
    assert body["total"] == 2
    row = next(u for u in body["items"] if u["role"] == "student")
    assert {"full_name", "username", "phone", "status", "selected_grade", "selected_track", "created_at", "last_active_at", "can_manage"} <= set(row)
    assert "password" not in str(row).lower().replace("must_change_password", "")
    assert client.get("/admin/users", params={"role": "admin"}, headers=admin_h).json()["total"] == 1
    assert client.get("/admin/users", params={"q": "abebe"}, headers=admin_h).json()["total"] == 1
    assert client.get("/admin/users", params={"q": "91234"}, headers=admin_h).json()["total"] == 1
    detail = client.get(f"/admin/users/{student['id']}", headers=admin_h).json()
    assert detail["progress"]["answered"] == 0


def test_students_cannot_reach_user_management(client, student):
    for method, path in [("GET", "/admin/users"), ("GET", "/admin/stats"), ("GET", "/admin/announcements"), ("GET", "/admin/reports"),
                         ("POST", f"/admin/users/{student['id']}/role"), ("POST", f"/admin/users/{student['id']}/suspend"),
                         ("DELETE", f"/admin/users/{student['id']}"), ("GET", "/admin/questions/import/template"), ("GET", "/admin/questions/import/history")]:
        assert client.request(method, path, headers=student["h"], json={} if method == "POST" else None).status_code == 403, (method, path)
        assert client.request(method, path).status_code == 401, (method, path)
    assert client.post("/admin/users", json={"role": "admin"}).status_code in (401, 405)       # no API creates an admin account directly


# ------------------------------------------------------------------ promotion chain
def test_promotion_chain_rules(client, admin_h, admin, grades):
    a_id, _ = mk(client, "Admin A", "0911111111", "admin_a")
    c_id, _ = mk(client, "Admin C", "0933333333", "admin_c")
    b_id, _ = mk(client, "Admin B", "0922222222", "admin_b")
    promote(client, admin_h, a_id)
    promote(client, admin_h, c_id)
    ah, _ = login(client, "admin_a")
    ch, _ = login(client, "admin_c")
    assert promote(client, ah, b_id)["promoted_by"]["id"] == a_id                  # A promotes B
    bh, _ = login(client, "admin_b")
    # B (set by A) cannot demote A; unrelated C cannot demote B; A can demote B; owner can demote anyone
    assert client.post(f"/admin/users/{a_id}/role", json={"role": "student"}, headers=bh).status_code == 403
    assert client.post(f"/admin/users/{b_id}/role", json={"role": "student"}, headers=ch).status_code == 403
    assert client.post(f"/admin/users/{b_id}/suspend", json={"reason": "x"}, headers=ch).status_code == 403
    assert client.post(f"/admin/users/{b_id}/reset-password", headers=ch).status_code == 403
    assert client.delete(f"/admin/users/{b_id}", params={"confirm": "true"}, headers=ch).status_code == 403
    assert client.post(f"/admin/users/{b_id}/role", json={"role": "student"}, headers=ah).json()["role"] == "student"
    assert client.post(f"/admin/users/{c_id}/role", json={"role": "student"}, headers=admin_h).json()["role"] == "student"


def test_chain_is_indirect_and_demotion_passes_children_up(client, admin_h, admin, grades):
    a_id, _ = mk(client, "Admin A", "0911111111", "admin_a")
    b_id, _ = mk(client, "Admin B", "0922222222", "admin_b")
    promote(client, admin_h, a_id)
    ah, _ = login(client, "admin_a")
    promote(client, ah, b_id)
    bh, _ = login(client, "admin_b")
    d_id, _ = mk(client, "Admin D", "0944444444", "admin_d")
    promote(client, bh, d_id)
    # A set B, B set D: A may manage D (indirect), D may not manage B or A
    dh, _ = login(client, "admin_d")
    assert client.get(f"/admin/users/{d_id}", headers=ah).json()["can_manage"] is True
    assert client.post(f"/admin/users/{b_id}/suspend", json={}, headers=dh).status_code == 403
    # owner demotes B: D now hangs under A (B's promoter)
    assert client.post(f"/admin/users/{b_id}/role", json={"role": "student"}, headers=admin_h).status_code == 200
    d = client.get(f"/admin/users/{d_id}", headers=admin_h).json()
    assert d["promoted_by"]["id"] == a_id and d["role"] == "admin"
    # B lost admin immediately: old token no longer opens admin routes, refresh tokens are revoked
    assert client.get("/admin/users", headers=bh).status_code == 403


def test_owner_and_self_protection(client, admin_h, admin, grades):
    a_id, _ = mk(client, "Admin A", "0911111111", "admin_a")
    promote(client, admin_h, a_id)
    ah, _ = login(client, "admin_a")
    owner_id = admin.id
    for path, kw in [(f"/admin/users/{owner_id}/role", {"json": {"role": "student"}}), (f"/admin/users/{owner_id}/suspend", {"json": {"reason": "x"}}),
                     (f"/admin/users/{owner_id}/reset-password", {}), (f"/admin/users/{owner_id}/logout-everywhere", {})]:
        assert client.post(path, headers=ah, **kw).status_code == 403, path
    assert client.delete(f"/admin/users/{owner_id}", params={"confirm": "true"}, headers=ah).status_code == 403
    # nobody can act on themselves, owner included
    assert client.post(f"/admin/users/{owner_id}/role", json={"role": "student"}, headers=admin_h).status_code == 403
    assert client.post(f"/admin/users/{owner_id}/suspend", json={}, headers=admin_h).status_code == 403
    assert client.delete(f"/admin/users/{owner_id}", params={"confirm": "true"}, headers=admin_h).status_code == 403
    assert client.post(f"/admin/users/{a_id}/suspend", json={}, headers=ah).status_code == 403
    assert client.post(f"/admin/users/{a_id}/role", json={"role": "student"}, headers=ah).status_code == 403


def test_promoting_does_not_create_owner_and_requires_active(client, admin_h, student):
    r = promote(client, admin_h, student["id"])
    assert r["role"] == "admin" and r["is_owner"] is False
    uid, _ = mk(client, "Susp Ended", "0955555555", "susp_x")
    client.post(f"/admin/users/{uid}/suspend", json={"reason": "r"}, headers=admin_h)
    assert client.post(f"/admin/users/{uid}/role", json={"role": "admin"}, headers=admin_h).status_code == 409


# ------------------------------------------------------------------ suspend
def test_suspend_blocks_next_request_signs_out_everywhere_and_shows_reason(client, admin_h, student):
    sid = student["id"]
    assert client.get("/me", headers=student["h"]).status_code == 200
    r = client.post(f"/admin/users/{sid}/suspend", json={"reason": "Cheating in exams", "days": 7}, headers=admin_h)
    assert r.status_code == 200 and r.json()["status"] == "suspended" and r.json()["suspended_reason"] == "Cheating in exams"
    assert client.get("/me", headers=student["h"]).status_code == 401                                   # next request refused
    assert client.post("/auth/refresh", json={"refresh_token": student["tokens"]["refresh_token"]}).status_code == 401
    lg = client.post("/auth/login", json={"identifier": "0912345678", "password": STUDENT_PW})
    assert lg.status_code == 403 and "Cheating in exams" in lg.json()["error"]["message"]
    assert lg.json()["error"]["until"]
    assert client.post(f"/admin/users/{sid}/activate", headers=admin_h).json()["status"] == "active"
    assert client.post("/auth/login", json={"identifier": "0912345678", "password": STUDENT_PW}).status_code == 200


def test_suspension_with_end_date_lifts_itself(client, admin_h, student, db):
    client.post(f"/admin/users/{student['id']}/suspend", json={"reason": "short"}, headers=admin_h)
    u = db.get(User, student["id"])
    u.suspended_until = utcnow() - timedelta(minutes=1)
    db.commit()
    assert client.post("/auth/login", json={"identifier": "0912345678", "password": STUDENT_PW}).status_code == 200
    assert client.get(f"/admin/users/{student['id']}", headers=admin_h).json()["status"] == "active"


def test_suspend_validation(client, admin_h, student):
    past = (utcnow() - timedelta(days=1)).isoformat()
    assert client.post(f"/admin/users/{student['id']}/suspend", json={"until": past}, headers=admin_h).status_code == 422
    assert client.post("/admin/users/99999/suspend", json={}, headers=admin_h).status_code == 404


def test_reset_password_forces_change_and_hides_old_password(client, admin_h, student):
    r = client.post(f"/admin/users/{student['id']}/reset-password", headers=admin_h)
    temp = r.json()["temporary_password"]
    assert len(temp) >= 6 and "password_hash" not in r.text
    assert client.post("/auth/login", json={"identifier": "0912345678", "password": STUDENT_PW}).status_code == 401
    h, body = login(client, "0912345678", temp)
    assert body["user"]["must_change_password"] is True
    assert client.get("/app/bank", headers=h).status_code == 403


def test_delete_user_needs_confirmation_and_is_audited(client, admin_h, student):
    assert client.delete(f"/admin/users/{student['id']}", headers=admin_h).status_code == 400
    assert client.delete(f"/admin/users/{student['id']}", params={"confirm": "true"}, headers=admin_h).status_code == 204
    assert client.post("/auth/login", json={"identifier": "0912345678", "password": STUDENT_PW}).status_code == 401
    acts = [a["action"] for a in client.get("/admin/audit", headers=admin_h).json()["items"]]
    assert "user.delete" in acts


def test_audit_records_role_and_suspension_changes(client, admin_h, student):
    promote(client, admin_h, student["id"])
    uid, _ = mk(client, "Other One", "0966666666", "other_one")
    client.post(f"/admin/users/{uid}/suspend", json={"reason": "r"}, headers=admin_h)
    client.post(f"/admin/users/{uid}/activate", headers=admin_h)
    acts = [a["action"] for a in client.get("/admin/audit", params={"size": 100}, headers=admin_h).json()["items"]]
    assert {"user.promote", "user.suspend", "user.activate"} <= set(acts)


# ------------------------------------------------------------------ tracks, archive, reorder
def test_subject_tracks_reach_the_app_and_student_can_choose_a_track(client, admin_h, student):
    make_content(client, admin_h, subject_key="physics", subject_name="Physics", title="U1")
    make_content(client, admin_h, subject_key="history", subject_name="History", title="U1")
    sid = {s["key"]: s["id"] for s in client.get("/admin/subjects", headers=admin_h).json()}
    assert client.patch(f"/admin/subjects/{sid['physics']}", json={"track": "natural"}, headers=admin_h).json()["track"] == "natural"
    assert client.patch(f"/admin/subjects/{sid['history']}", json={"track": "social"}, headers=admin_h).json()["track"] == "social"
    assert client.patch(f"/admin/subjects/{sid['history']}", json={"track": "mars"}, headers=admin_h).status_code == 422
    subs = {s["id"]: s["t"] for s in client.get("/app/bank", headers=student["h"]).json()["subjects"]}
    assert subs == {"physics": "natural", "history": "social"}
    me = client.patch("/me", json={"selected_grade": 12, "selected_track": "natural"}, headers=student["h"]).json()
    assert me["selected_track"] == "natural" and me["selected_grade"] == 12
    assert client.patch("/me", json={"selected_track": "all"}, headers=student["h"]).json()["selected_track"] is None
    assert client.patch("/me", json={"selected_track": "arts"}, headers=student["h"]).status_code == 422


def test_archiving_hides_from_students_but_keeps_everything(client, admin_h, student):
    c = make_content(client, admin_h, n=3, title="Arch unit")
    qk = client.get("/app/bank", headers=student["h"]).json()["units"][0]["qs"][0]["k"]
    assert client.delete(f"/admin/units/{c['unit_id']}", headers=admin_h).json()["archived"] is True
    assert client.get("/app/bank", headers=student["h"]).json()["units"] == []
    assert c["unit_id"] not in [u["id"] for u in client.get("/admin/units", headers=admin_h).json()]
    assert c["unit_id"] in [u["id"] for u in client.get("/admin/units", params={"include_archived": "true"}, headers=admin_h).json()]
    assert client.get(f"/admin/questions/{c['qids'][0]}", headers=admin_h).json()["deleted"] is False        # questions untouched
    client.post(f"/admin/units/{c['unit_id']}/restore", headers=admin_h)
    assert client.get("/app/bank", headers=student["h"]).json()["units"][0]["qs"][0]["k"] == qk
    # subject and grade archive too
    client.delete(f"/admin/subjects/{c['subject_id']}", headers=admin_h)
    assert client.get("/app/bank", headers=student["h"]).json()["units"] == []
    client.post(f"/admin/subjects/{c['subject_id']}/restore", headers=admin_h)
    client.delete(f"/admin/grades/{c['grade_id']}", headers=admin_h)
    assert client.get("/app/bank", headers=student["h"]).json()["units"] == []
    client.post(f"/admin/grades/{c['grade_id']}/restore", headers=admin_h)
    assert len(client.get("/app/bank", headers=student["h"]).json()["units"]) == 1


def test_reorder_subjects_and_units(client, admin_h, student):
    make_content(client, admin_h, subject_key="math", subject_name="Mathematics", title="U1")
    make_content(client, admin_h, subject_key="physics", subject_name="Physics", title="U1")
    ids = [s["id"] for s in client.get("/admin/subjects", headers=admin_h).json()]
    r = client.post("/admin/subjects/reorder", json={"subject_ids": list(reversed(ids))}, headers=admin_h)
    assert r.status_code == 200
    assert [s["id"] for s in client.get("/app/bank", headers=student["h"]).json()["subjects"]] == ["physics", "math"]
    assert client.post("/admin/subjects/reorder", json={"subject_ids": [ids[0], ids[0]]}, headers=admin_h).status_code == 422


# ------------------------------------------------------------------ excel import
def xlsx(rows, header=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "Questions"
    ws.append(header or ["grade", "subject", "unit", "question", "option_a", "option_b", "option_c", "option_d", "correct_answer", "explanation", "image", "difficulty"])
    for r in rows:
        ws.append(r)
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def up(client, h, data, kind="validate", **form):
    return client.post(f"/admin/questions/import/{kind}", files={"file": ("q.xlsx", data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}, data=form, headers=h)


ROWS = [[11, "Mathematics", "Unit 1", "What is 12 x 12?", "122", "144", "124", "148", "B", "12 squared", "", "easy"],
        [11, "Mathematics", "Unit 1", "Bad row", "1", "2", "3", "4", "Z", "", "", ""],
        [11, "Mathematics", "1", "What is 5 + 5?", "10", "11", "12", "13", "10", "", "", "medium"]]


def test_template_is_a_valid_workbook_with_the_fixed_columns(client, admin_h):
    make_content(client, admin_h, title="Unit 1")
    r = client.get("/admin/questions/import/template", headers=admin_h)
    assert r.status_code == 200 and r.content[:2] == b"PK"
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Questions", "Instructions", "Reference"]
    assert [c.value for c in wb["Questions"][1]] == ["grade", "subject", "unit", "question", "option_a", "option_b", "option_c", "option_d", "correct_answer", "explanation", "image", "difficulty"]
    assert any(row[1].value == "math" for row in wb["Reference"].iter_rows(min_row=2))
    assert up(client, admin_h, r.content).status_code == 200                       # the untouched template itself is importable


def test_excel_preview_reports_errors_and_duplicates_then_imports_drafts(client, admin_h, student):
    make_content(client, admin_h, n=1, title="Unit 1")
    v = up(client, admin_h, xlsx(ROWS)).json()
    assert v["summary"] == {"total": 3, "ok": 2, "errors": 1, "duplicates": 0}
    assert "does not match" in v["rows"][1]["errors"][0]
    c = up(client, admin_h, xlsx(ROWS), "commit")
    assert c.status_code == 200 and c.json()["created"]["created"] == 2
    bank = client.get("/app/bank", headers=student["h"]).json()
    assert [q["q"] for u in bank["units"] for q in u["qs"] if "x 12" in q["q"] or "5 + 5" in q["q"]] == []          # imported as drafts: students see nothing
    qs = client.get("/admin/questions", params={"status": "draft"}, headers=admin_h).json()
    assert qs["total"] == 2
    again = up(client, admin_h, xlsx(ROWS)).json()
    assert again["summary"]["duplicates"] == 2


def test_excel_duplicate_policies_and_undo(client, admin_h):
    make_content(client, admin_h, n=1, title="Unit 1")
    up(client, admin_h, xlsx(ROWS[:1]), "commit")
    changed = [[11, "Mathematics", "Unit 1", "What is 12 x 12?", "122", "144", "124", "148", "B", "NEW explanation", "", "hard"]]
    skip = up(client, admin_h, xlsx(changed), "commit", policy="skip").json()
    assert skip["created"]["skipped_duplicates"] == 1 and skip["created"]["created"] == 0
    upd = up(client, admin_h, xlsx(changed), "commit", policy="update").json()
    assert upd["created"]["updated"] == 1
    q = next(x for x in client.get("/admin/questions", params={"q": "12 x 12"}, headers=admin_h).json()["items"])
    assert q["explanation"] == "NEW explanation" and q["difficulty"] == "hard" and q["version"] == 2
    add = up(client, admin_h, xlsx(changed), "commit", policy="add").json()
    assert add["created"]["created"] == 1
    assert client.get("/admin/questions", params={"q": "12 x 12"}, headers=admin_h).json()["total"] == 2
    hist = client.get("/admin/questions/import/history", headers=admin_h).json()
    assert hist["total"] == 4 and hist["items"][0]["filename"] == "q.xlsx" and hist["items"][0]["admin"] == "Site Admin"
    # undo the update batch: explanation restored; undo the add batch: the copy disappears
    assert client.post(f"/admin/questions/import/{upd['batch_id']}/undo", headers=admin_h).json() == {"removed": 0, "restored": 1}
    q = next(x for x in client.get("/admin/questions", params={"q": "12 x 12"}, headers=admin_h).json()["items"] if x["id"] == q["id"])
    assert q["explanation"] == "12 squared"
    assert client.post(f"/admin/questions/import/{add['batch_id']}/undo", headers=admin_h).json()["removed"] == 1
    assert client.get("/admin/questions", params={"q": "12 x 12"}, headers=admin_h).json()["total"] == 1
    assert client.post(f"/admin/questions/import/{add['batch_id']}/undo", headers=admin_h).status_code == 409


def test_undo_of_a_batch_that_created_units_and_subjects(client, admin_h, student):
    make_content(client, admin_h, n=1, title="Unit 1")
    rows = [[12, "Brand New Subject", "Unit 1", "Q one here?", "a1", "b1", "c1", "d1", "C", "", "", ""]]
    c = up(client, admin_h, xlsx(rows), "commit", create_missing="true", publish="true").json()
    assert c["created"]["subjects"] == 1 and c["created"]["units"] == 1
    assert any(s["name"] == "Brand New Subject" for s in client.get("/admin/subjects", headers=admin_h).json())
    client.post(f"/admin/questions/import/{c['batch_id']}/undo", headers=admin_h)
    assert not any(s["name"] == "Brand New Subject" for s in client.get("/admin/subjects", headers=admin_h).json())
    assert all("Q one here" not in q["q"] for u in client.get("/app/bank", headers=student["h"]).json()["units"] for q in u["qs"])


def test_excel_rejects_bad_files_and_neutralises_formulas(client, admin_h):
    assert up(client, admin_h, b"PK\x03\x04 not really a workbook").status_code == 422
    assert up(client, admin_h, xlsx([], header=["a", "b"])).status_code == 422
    make_content(client, admin_h, n=0, title="Unit 1")
    bad_img = [[11, "Mathematics", "Unit 1", "Has an outside image?", "1", "2", "3", "4", "A", "", "http://evil.example/x.png", ""]]
    assert up(client, admin_h, xlsx(bad_img)).json()["summary"]["errors"] == 1
    ok_img = [[11, "Mathematics", "Unit 1", "Has an uploaded image?", "1", "2", "3", "4", "A", "", "/media/abc.png", ""]]
    assert up(client, admin_h, xlsx(ok_img)).json()["summary"]["ok"] == 1
    r = client.post("/admin/questions", headers=admin_h, json={"unit_id": client.get("/admin/units", headers=admin_h).json()[0]["id"], "text": "=SUM(1,1) is a formula?",
                                                                "options": [{"text": "+cmd", "is_correct": True}, {"text": "no"}]})
    assert r.status_code == 201
    wb = load_workbook(io.BytesIO(client.get("/admin/questions/export", params={"format": "xlsx"}, headers=admin_h).content))
    cells = [c.value for row in wb["Questions"].iter_rows(min_row=2) for c in row if isinstance(c.value, str)]
    assert "'=SUM(1,1) is a formula?" in cells and "'+cmd" in cells


def test_export_then_import_round_trip_recognises_everything_as_existing(client, admin_h):
    make_content(client, admin_h, n=3, title="Unit 1")
    data = client.get("/admin/questions/export", params={"format": "xlsx"}, headers=admin_h).content
    v = up(client, admin_h, data).json()
    assert v["summary"]["total"] == 3 and v["summary"]["duplicates"] == 3 and v["summary"]["errors"] == 0


# ------------------------------------------------------------------ statistics
def push_state(client, h, answers, history=()):
    st = {"b": {}, "f": {}, "a": {k: [ts, ok] for k, ts, ok in answers}, "h": list(history)}
    r = client.put("/app/state", json={"state": st}, headers=h)
    assert r.status_code == 200, r.text


def test_statistics_are_computed_from_student_data_and_can_be_exported(client, admin_h, student):
    c = make_content(client, admin_h, n=3, title="Unit 1")
    texts = [client.get(f"/admin/questions/{q}", headers=admin_h).json()["text"] for q in c["qids"]]
    keys = [qkey(t) for t in texts]
    now = int(time.time() * 1000)
    push_state(client, student["h"], [(keys[0], now, 1), (keys[1], now, 0), (keys[2], now, 1)],
               [{"id": "s1", "ti": "Maths", "m": "practice", "n": 3, "c": 2, "ts": now, "d": 120}, {"id": "s2", "ti": "Mock", "m": "exam", "n": 10, "c": 5, "ts": now, "d": 600}])
    _, other = None, mk(client, "Second Student", "0944444444", "second_s")[1]
    push_state(client, other, [(keys[0], now, 0), (keys[1], now, 0)])
    s = client.get("/admin/stats", params={"min_answers": 2}, headers=admin_h).json()
    assert s["users"]["total"] == 3 and s["users"]["students"] == 2 and s["users"]["admins"] == 1
    assert s["learning"]["questions_answered"] == 5 and s["learning"]["correct"] == 2 and s["learning"]["accuracy"] == 40
    assert s["learning"]["active_learners"] == 2
    assert s["learning"]["by_subject"][0] == {"subject": "Mathematics", "answered": 5, "correct": 2, "accuracy": 40}
    assert s["learning"]["by_grade"][0]["grade"] == 11
    assert s["learning"]["sessions"]["total"] == 2
    assert s["learning"]["sessions"]["by_mode"]["exam"]["average_percent"] == 50 and s["learning"]["sessions"]["by_mode"]["practice"]["average_minutes"] == 2.0
    assert sum(h["sessions"] for h in s["learning"]["busiest_hours"]) == 2
    hardest = s["questions"]["hardest"]
    assert hardest[0]["accuracy"] == 0 and hardest[0]["answered"] == 2                         # keys[1]: wrong twice
    assert s["questions"]["never_answered_count"] == 0
    assert s["content"]["questions"]["published"] == 3
    d = client.get(f"/admin/users/{student['id']}", headers=admin_h).json()["progress"]
    assert d["answered"] == 3 and d["correct"] == 2 and d["sessions"] == 2 and d["exams"] == 1
    # date range: a window in the past contains none of it
    old = client.get("/admin/stats", params={"from": "2020-01-01", "to": "2020-01-31"}, headers=admin_h).json()
    assert old["learning"]["questions_answered"] == 0 and old["learning"]["sessions"]["total"] == 0
    assert client.get("/admin/stats", params={"from": "2021-01-02", "to": "2021-01-01"}, headers=admin_h).status_code == 422
    # exports
    x = client.get("/admin/stats/export", params={"format": "xlsx"}, headers=admin_h)
    wb = load_workbook(io.BytesIO(x.content))
    assert "accuracy_by_subject" in wb.sheetnames and "summary" in wb.sheetnames
    csv = client.get("/admin/stats/export", params={"format": "csv", "table": "accuracy_by_subject"}, headers=admin_h)
    assert csv.status_code == 200 and "Mathematics,5,2,40" in csv.text
    assert client.get("/admin/stats/export", params={"format": "csv", "table": "nope"}, headers=admin_h).status_code == 404


def test_stats_flag_questions_nobody_answers_and_suspicious_answer_keys(client, admin_h, student):
    c = make_content(client, admin_h, n=2, title="Unit 1")
    k = qkey(client.get(f"/admin/questions/{c['qids'][0]}", headers=admin_h).json()["text"])
    now = int(time.time() * 1000)
    push_state(client, student["h"], [(k, now, 0)])
    s = client.get("/admin/stats", params={"min_answers": 1}, headers=admin_h).json()
    assert s["questions"]["never_answered_count"] == 1
    assert [x["accuracy"] for x in s["questions"]["may_have_wrong_answer_key"]] == [0]


# ------------------------------------------------------------------ announcements and reports
def test_announcements_reach_students_until_expired_or_switched_off(client, admin_h, student):
    a = client.post("/admin/announcements", json={"title": "Mock exam Friday", "body": "Bring a pen."}, headers=admin_h).json()
    exp = client.post("/admin/announcements", json={"title": "Old news", "body": "x", "expires_at": (utcnow() - timedelta(days=1)).isoformat()}, headers=admin_h).json()
    assert [x["title"] for x in client.get("/app/announcements", headers=student["h"]).json()] == ["Mock exam Friday"]
    client.patch(f"/admin/announcements/{a['id']}", json={"active": False}, headers=admin_h)
    assert client.get("/app/announcements", headers=student["h"]).json() == []
    assert client.delete(f"/admin/announcements/{exp['id']}", headers=admin_h).status_code == 204
    assert client.get("/app/announcements").status_code == 401


def test_students_report_questions_and_admins_work_the_queue(client, admin_h, student):
    c = make_content(client, admin_h, n=1, title="Unit 1")
    key = client.get("/app/bank", headers=student["h"]).json()["units"][0]["qs"][0]["k"]
    r = client.post("/app/reports", json={"key": key, "reason": "wrong_answer", "note": "The answer should be 1"}, headers=student["h"])
    assert r.status_code == 201 and r.json()["duplicate"] is False
    assert client.post("/app/reports", json={"key": key, "reason": "typo"}, headers=student["h"]).json()["duplicate"] is True
    assert client.post("/app/reports", json={"key": "z" * 64, "reason": "typo"}, headers=student["h"]).status_code == 404
    assert client.post("/app/reports", json={"key": key, "reason": "rude"}, headers=student["h"]).status_code == 422
    q = client.get("/admin/reports", headers=admin_h).json()
    assert q["total"] == 1 and q["items"][0]["question_id"] == c["qids"][0] and q["items"][0]["reporter"] == "Abebe Kebede"
    rid = q["items"][0]["id"]
    assert client.post(f"/admin/reports/{rid}/resolve", headers=admin_h).json()["status"] == "resolved"
    assert client.get("/admin/reports", headers=admin_h).json()["total"] == 0
    assert client.get("/admin/reports", params={"status": "resolved"}, headers=admin_h).json()["total"] == 1


def test_exam_date_setting_is_public_config(client, admin_h):
    assert client.put("/admin/settings", json={"exam_name": "Ethiopian University Entrance Exam", "exam_date": "2027-07-05"}, headers=admin_h).status_code == 200
    cfg = client.get("/config").json()
    assert cfg["exam_date"] == "2027-07-05" and cfg["exam_name"].startswith("Ethiopian")
    assert client.put("/admin/settings", json={"exam_date": "next week"}, headers=admin_h).status_code == 422
